"""Рисунки 1, 3 и 4 статьи: таблицы и нативные диаграммы Excel из файлов results/.

Скрипт ничего не считает заново: он берёт готовые результаты diagnostics.py и train_models.py
и записывает article/figures.xlsx. Рисунок 2 (схема пайплайна) лежит в
article/figure2_pipeline.drawio и правится в draw.io.
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.drawing.text import Font as ChartFont
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.worksheet import Worksheet

from common import RESULTS_DIR, ROOT

ARTICLE_DIR = ROOT / "article"
OUTPUT_PATH = ARTICLE_DIR / "figures.xlsx"
MONTHLY_PATH = RESULTS_DIR / "monthly_medians.csv"
IMPORTANCE_PATH = RESULTS_DIR / "permutation_importance.csv"
TIME_PATH = RESULTS_DIR / "high_views_by_time.csv"

BLUE, ORANGE, INK, INK_SECONDARY, GRID = "2A78D6", "EB6834", "0B0B0B", "52514E", "D9D9D9"
FONT_NAME = "Times New Roman"
TOP_FEATURES = 5
# Часы 00–05 не показаны на рисунке 4: в части ячеек меньше 30 постов (у блогеров в 04:00 их 3).
FIRST_PLOTTED_HOUR = 6
FEATURE_LABELS = {
    "text": "Text (100 TF-IDF/SVD components)",
    "weekday": "Weekday",
    "hour": "Hour of publication",
    "n_emoji": "Number of emoji",
    "actor_type": "Layer",
}
LAYER_LABELS = {"government": "Government layer", "blogger": "Blogger layer"}
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=11)
CELL_FONT = Font(name=FONT_NAME, size=11)
NOTE_FONT = Font(name=FONT_NAME, size=10, italic=True, color=INK_SECONDARY)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RULE = Side(style="thin", color="A6A6A6")


def chart_text(size: int = 1000, color: str = INK_SECONDARY) -> RichText:
    props = CharacterProperties(latin=ChartFont(typeface=FONT_NAME), sz=size, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), endParaRPr=props)])


def style_axes(chart: LineChart | BarChart, x_title: str | None, y_title: str) -> None:
    for axis, title in ((chart.x_axis, x_title), (chart.y_axis, y_title)):
        axis.delete = False
        axis.title = title
        axis.txPr = chart_text()
        if title is not None:
            title_font = ChartFont(typeface=FONT_NAME)
            axis.title.tx.rich.p[0].pPr = ParagraphProperties(
                defRPr=CharacterProperties(latin=title_font, sz=1000, solidFill=INK)
            )
        axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="A6A6A6", w=9525))
        if axis.numFmt is not None:
            axis.numFmt.sourceLinked = False
    grid_line = LineProperties(solidFill=GRID, w=6350)
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=grid_line)
    chart.x_axis.majorGridlines = None
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))


def style_line(series, color: str, symbol: str) -> None:
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = 22225
    series.marker.symbol = symbol
    series.marker.size = 6
    series.marker.graphicalProperties = GraphicalProperties(
        solidFill=color, ln=LineProperties(solidFill=color)
    )
    series.smooth = False


def write_table(sheet: Worksheet, title: str, header: list[str], rows: list[list]) -> int:
    sheet["A1"] = title
    sheet["A1"].font = Font(name=FONT_NAME, bold=True, size=12)
    for column, name in enumerate(header, start=1):
        cell = sheet.cell(row=3, column=column, value=name)
        cell.font, cell.alignment = HEADER_FONT, CENTER
        cell.border = Border(top=RULE, bottom=RULE)
        sheet.column_dimensions[cell.column_letter].width = 15
    for offset, values in enumerate(rows, start=4):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=offset, column=column, value=value)
            cell.font, cell.alignment = CELL_FONT, CENTER
    return 3 + len(rows)


def add_note(sheet: Worksheet, row: int, text: str) -> None:
    sheet.cell(row=row, column=1, value=text).font = NOTE_FONT


def source_note(path: Path, script: str) -> str:
    return f"Source: {path.relative_to(ROOT).as_posix()} ({script})."


def month_label(month: str) -> str:
    return pd.Timestamp(f"{month}-01").strftime("%b %Y")


def figure_1(sheet: Worksheet) -> LineChart:
    medians = pd.read_csv(MONTHLY_PATH)
    wide = medians.pivot(index="month", columns="layer").sort_index()
    rows = [
        [
            month_label(month),
            wide.loc[month, ("median_views", "government")],
            wide.loc[month, ("median_views", "blogger")],
            f"=C{row}/B{row}",
            int(wide.loc[month, ("n", "government")]),
            int(wide.loc[month, ("n", "blogger")]),
            round(wide.loc[month, ("median_age_days", "all")]),
        ]
        for row, month in enumerate(wide.index, start=4)
    ]
    header = [
        "Month",
        "Government, median views",
        "Blogger, median views",
        "Blogger ÷ government",
        "Government, posts",
        "Blogger, posts",
        "Median age at collection, days (all posts)",
    ]
    last = write_table(sheet, "Figure 1. Median views of a post by month and layer", header, rows)
    for row in range(4, last + 1):
        sheet.cell(row=row, column=4).number_format = "0.00"
    add_note(sheet, last + 2, source_note(MONTHLY_PATH, "diagnostics.py"))
    chart = LineChart()
    chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=4, max_row=last))
    chart.set_categories(Reference(sheet, min_col=1, min_row=4, max_row=last))
    chart.series[0].tx = SeriesLabel(v=LAYER_LABELS["government"])
    chart.series[1].tx = SeriesLabel(v=LAYER_LABELS["blogger"])
    style_line(chart.series[0], BLUE, "circle")
    style_line(chart.series[1], ORANGE, "square")
    style_axes(chart, "Month of publication", "Median views of a post")
    chart.y_axis.scaling.min, chart.y_axis.scaling.max, chart.y_axis.majorUnit = 0, 3000, 500
    chart.y_axis.number_format = "#,##0"
    chart.legend.position, chart.legend.txPr = "t", chart_text(color=INK)
    chart.width, chart.height = 16, 9.5
    return chart


def figure_3(sheet: Worksheet) -> BarChart:
    importance = pd.read_csv(IMPORTANCE_PATH).nlargest(TOP_FEATURES, "importance_mean")
    rows = [
        [FEATURE_LABELS.get(feature, feature), round(value, 4)]
        for feature, value in zip(importance["feature"], importance["importance_mean"], strict=True)
    ]
    header = ["Feature", "Mean drop in test ROC-AUC"]
    last = write_table(sheet, "Figure 3. Permutation importance, five largest", header, rows)
    sheet.column_dimensions["A"].width = 36
    add_note(sheet, last + 2, source_note(IMPORTANCE_PATH, "train_models.py"))
    chart = BarChart()
    chart.type, chart.gapWidth = "bar", 60
    chart.add_data(Reference(sheet, min_col=2, min_row=3, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=4, max_row=last))
    bars = chart.series[0]
    bars.graphicalProperties = GraphicalProperties(solidFill=BLUE, ln=LineProperties(noFill=True))
    bars.dLbls = DataLabelList()
    bars.dLbls.showVal = True
    bars.dLbls.showSerName = bars.dLbls.showCatName = bars.dLbls.showLegendKey = False
    bars.dLbls.numFmt, bars.dLbls.position = "0.0000", "outEnd"
    bars.dLbls.txPr = chart_text(color=INK)
    style_axes(chart, None, "Mean drop in test ROC-AUC after permutation")
    chart.x_axis.scaling.orientation = "maxMin"
    chart.y_axis.crosses = "max"
    chart.y_axis.scaling.min, chart.y_axis.scaling.max, chart.y_axis.majorUnit = 0, 0.09, 0.01
    chart.y_axis.number_format = "0.00"
    chart.legend = None
    chart.width, chart.height = 16, 7.5
    return chart


def figure_4(sheet: Worksheet) -> LineChart:
    time = pd.read_csv(TIME_PATH)
    hours = time[time["variable"] == "hour"].pivot(index="value", columns="layer").sort_index()
    rows = [
        [
            f"{hour:02d}",
            hours.loc[hour, ("share", "government")],
            hours.loc[hour, ("share", "blogger")],
            int(hours.loc[hour, ("n", "government")]),
            int(hours.loc[hour, ("n", "blogger")]),
        ]
        for hour in hours.index
        if hour >= FIRST_PLOTTED_HOUR
    ]
    header = [
        "Hour (UTC+5)",
        "Government, share of class 1",
        "Blogger, share of class 1",
        "Government, posts",
        "Blogger, posts",
    ]
    last = write_table(sheet, "Figure 4. Share of class 1 by hour and layer", header, rows)
    for row in range(4, last + 1):
        for column in (2, 3):
            sheet.cell(row=row, column=column).number_format = "0.0%"
    add_note(sheet, last + 2, source_note(TIME_PATH, "diagnostics.py"))
    add_note(sheet, last + 3, "Hours 00–05 are left out: some cells hold fewer than 30 posts.")
    chart = LineChart()
    chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=4, max_row=last))
    chart.set_categories(Reference(sheet, min_col=1, min_row=4, max_row=last))
    chart.series[0].tx = SeriesLabel(v=LAYER_LABELS["government"])
    chart.series[1].tx = SeriesLabel(v=LAYER_LABELS["blogger"])
    style_line(chart.series[0], BLUE, "circle")
    style_line(chart.series[1], ORANGE, "square")
    style_axes(chart, "Hour of publication (UTC+5)", "Share of class 1")
    chart.y_axis.scaling.min, chart.y_axis.scaling.max, chart.y_axis.majorUnit = 0.3, 0.8, 0.1
    chart.y_axis.number_format = "0%"
    chart.legend.position, chart.legend.txPr = "t", chart_text(color=INK)
    chart.width, chart.height = 16, 9.5
    return chart


def main() -> None:
    ARTICLE_DIR.mkdir(exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, build, anchor in (
        ("Figure 1", figure_1, "I3"),
        ("Figure 3", figure_3, "D3"),
        ("Figure 4", figure_4, "G3"),
    ):
        sheet = workbook.create_sheet(name)
        sheet.add_chart(build(sheet), anchor)
    workbook.save(OUTPUT_PATH)
    print(f"Записано: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
